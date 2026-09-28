from pathlib import Path
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "hospital_management.xlsx"

USER_HEADERS = ["Username", "Password", "Full Name", "Email"]

PATIENT_HEADERS = [
    "Patient ID", "Name", "Age", "Gender",
    "Blood Group", "Phone", "Address", "Disease"
]

APPOINTMENT_HEADERS = [
    "Appointment ID", "Patient ID", "Patient Name",
    "Doctor", "Date", "Time", "Status"
]


def create_excel():
    """Create the database workbook and required sheets if they are missing."""
    if not EXCEL_FILE.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Users"
        ws.append(USER_HEADERS)

        ws = wb.create_sheet("Patients")
        ws.append(PATIENT_HEADERS)

        ws = wb.create_sheet("Appointments")
        ws.append(APPOINTMENT_HEADERS)

        wb.save(EXCEL_FILE)
        wb.close()
        return

    wb = load_workbook(EXCEL_FILE)

    required = {
        "Users": USER_HEADERS,
        "Patients": PATIENT_HEADERS,
        "Appointments": APPOINTMENT_HEADERS,
    }

    changed = False
    for sheet_name, headers in required.items():
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(sheet_name)
            ws.append(headers)
            changed = True
        else:
            ws = wb[sheet_name]
            # If an empty sheet exists, add its header row.
            if ws.max_row == 1 and all(v is None for v in ws[1]):
                ws.delete_rows(1)
                ws.append(headers)
                changed = True

    if changed:
        wb.save(EXCEL_FILE)

    wb.close()


def get_records(sheet_name):
    create_excel()

    wb = load_workbook(EXCEL_FILE, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return []

        ws = wb[sheet_name]
        return list(ws.iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()


def add_record(sheet_name, values):
    create_excel()

    wb = load_workbook(EXCEL_FILE)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Unknown sheet: {sheet_name}")

        wb[sheet_name].append(list(values))
        wb.save(EXCEL_FILE)
    finally:
        wb.close()


def update_record(sheet_name, id_column, record_id, values):
    create_excel()

    wb = load_workbook(EXCEL_FILE)
    try:
        if sheet_name not in wb.sheetnames:
            return False

        ws = wb[sheet_name]

        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row, id_column).value).strip() == str(record_id).strip():
                for column, value in enumerate(values, start=1):
                    ws.cell(row, column).value = value

                wb.save(EXCEL_FILE)
                return True

        return False
    finally:
        wb.close()


def delete_record(sheet_name, id_column, record_id):
    create_excel()

    wb = load_workbook(EXCEL_FILE)
    try:
        if sheet_name not in wb.sheetnames:
            return False

        ws = wb[sheet_name]

        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row, id_column).value).strip() == str(record_id).strip():
                ws.delete_rows(row, 1)
                wb.save(EXCEL_FILE)
                return True

        return False
    finally:
        wb.close()


def username_exists(username):
    username = str(username).strip().lower()

    return any(
        str(row[0]).strip().lower() == username
        for row in get_records("Users")
        if row and row[0] is not None
    )


def register_user(username, password, full_name, email):
    if username_exists(username):
        return False

    add_record(
        "Users",
        [str(username).strip(), password, str(full_name).strip(), str(email).strip()]
    )
    return True


def validate_login(username, password):
    username = str(username).strip()

    return any(
        str(row[0]).strip() == username and str(row[1]) == str(password)
        for row in get_records("Users")
        if row and row[0] is not None and row[1] is not None
    )
